from datetime import datetime
from pathlib import Path

from docx import Document
from langchain_text_splitters import RecursiveCharacterTextSplitter
from pypdf import PdfReader

from config import settings


class DocumentLoadError(Exception):
    pass


class UnsupportedDocumentError(DocumentLoadError):
    pass


def validate_supported_file(filename: str) -> str:
    suffix = Path(filename).suffix.lower()
    if suffix not in settings.allowed_extensions:
        allowed = ", ".join(sorted(settings.allowed_extensions))
        raise UnsupportedDocumentError(f"Unsupported file type. Allowed types: {allowed}.")
    return suffix


def extract_text(file_path: Path) -> str:
    return "\n".join(part["text"] for part in extract_parts(file_path))


def extract_parts(file_path: Path) -> list[dict]:
    suffix = validate_supported_file(file_path.name)
    try:
        if suffix == settings.pdf_extension:
            return _extract_pdf_parts(file_path)
        if suffix == settings.docx_extension:
            return _extract_docx_parts(file_path)
        if suffix == settings.txt_extension:
            return _extract_txt_parts(file_path)
    except UnsupportedDocumentError:
        raise
    except Exception as exc:
        raise DocumentLoadError(f"Could not read document '{file_path.name}'.") from exc

    raise UnsupportedDocumentError(f"Unsupported file type: {suffix}")


def chunk_text(text: str, filename: str) -> list[dict]:
    return chunk_parts([{"text": text, "page": None}], filename)


def chunk_parts(parts: list[dict], filename: str) -> list[dict]:
    splitter = RecursiveCharacterTextSplitter(
        chunk_size=settings.chunk_size,
        chunk_overlap=settings.chunk_overlap,
        separators=["\n\n", "\n", ". ", " ", ""],
    )
    chunks: list[dict] = []
    timestamp = datetime.utcnow().isoformat()

    for part in parts:
        clean_text = " ".join(str(part.get("text", "")).split())
        if not clean_text:
            continue
        for text in splitter.split_text(clean_text):
            if not text.strip():
                continue
            chunk_index = len(chunks)
            page = part.get("page")
            chunks.append(
                {
                    "id": f"{filename}:{chunk_index}",
                    "text": text,
                    "metadata": {
                        "filename": filename,
                        "document_name": filename,
                        "source_file": filename,
                        "page": page if page is not None else "",
                        "chunk": chunk_index,
                        "chunk_id": f"{filename}:{chunk_index}",
                        "timestamp": timestamp,
                    },
                }
            )

    if not chunks:
        raise DocumentLoadError("The uploaded document does not contain readable text.")
    return chunks


def _extract_pdf_parts(file_path: Path) -> list[dict]:
    reader = PdfReader(str(file_path))
    return [
        {"text": page.extract_text() or "", "page": index + 1}
        for index, page in enumerate(reader.pages)
    ]


def _extract_docx_parts(file_path: Path) -> list[dict]:
    document = Document(str(file_path))
    paragraphs = [paragraph.text for paragraph in document.paragraphs]
    table_cells = [
        cell.text
        for table in document.tables
        for row in table.rows
        for cell in row.cells
    ]
    return [{"text": "\n".join(paragraphs + table_cells), "page": None}]


def _extract_txt_parts(file_path: Path) -> list[dict]:
    data = file_path.read_bytes()
    for encoding in ("utf-8", "utf-8-sig", "cp1252", "latin-1"):
        try:
            return [{"text": data.decode(encoding), "page": None}]
        except UnicodeDecodeError:
            continue
    raise DocumentLoadError("Could not decode the text file.")
